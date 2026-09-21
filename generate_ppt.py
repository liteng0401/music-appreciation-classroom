# -*- coding: utf-8 -*-
"""
为《高中音乐鉴赏》每节生成一份上课用 PPT 课件（.pptx）。
板块：知识点讲解 + 变易教学设计 + 作品鉴赏 + 随堂练习。
变易设计按"变易理论"（审辨 + 对比/类合/区分/融合）为每单元新写，结合该单元真实作品与概念。
用法：python3 generate_ppt.py  ->  输出到 ./ppt/
"""
import os, re, json, glob

ROOT = "/Users/lt/WorkBuddy/2026-09-21-01-46-33/music-appreciation-web"
SKILL = "/Users/lt/.workbuddy/skills/music-appreciation"
LOGO = os.path.join(ROOT, "school-logo.png")
NAMEIMG = os.path.join(ROOT, "school-name.png")
OUT = os.path.join(ROOT, "ppt")
os.makedirs(OUT, exist_ok=True)

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BRAND = RGBColor(0xC0, 0x39, 0x2B)
BRAND2 = RGBColor(0x16, 0x63, 0x8A)
INK = RGBColor(0x23, 0x30, 0x3A)
MUTED = RGBColor(0x6B, 0x77, 0x80)
PANEL = RGBColor(0xF4, 0xF1, 0xEA)
ACCENT = RGBColor(0xF0, 0xB4, 0x29)

CN_NUM = "零一二三四五六七八九十"
def cn_to_int(s):
    # 支持 一~十九 简单映射
    s = s.strip()
    if s in CN_NUM: return CN_NUM.index(s)
    if len(s) == 2:
        if s[0] == "十": return 10 + (CN_NUM.index(s[1]) if s[1] in CN_NUM else 0)
        if s[1] == "十": return CN_NUM.index(s[0]) * 10
    if len(s) == 3 and s[1] == "十":
        return CN_NUM.index(s[0]) * 10 + CN_NUM.index(s[2])
    if s == "十": return 10
    return 0

def group_of(uid):
    n = int(uid[2:4])
    return "序篇" if n == 0 else ("上篇" if n <= 8 else "下篇")

# ---------- 解析章节 ----------
def parse_section_names(text):
    """返回 [(绝对节序号:int, 节名:str), ...]，兼容两种格式：
       上篇: ## 第一节 X　|　第二节 Y
       下篇: # 第九单元 名（第十七节 A · 第十八节 B）"""
    line = ""
    for ln in text.splitlines():
        if "节" in ln and (ln.startswith("# ") or ln.startswith("## ")):
            line = ln; break
    # 去掉 # 号与单元名前缀
    line = re.sub(r'^#+\s*', '', line)
    line = re.sub(r'第[一二三四五六七八九十]+单元[^（(]*', '', line)  # 去单元名
    line = line.strip("（()） ")
    pat = re.compile(r'第([一二三四五六七八九十]+)节\s*([^|｜·\n（）()]+)')
    out = []
    for m in pat.finditer(line):
        out.append((cn_to_int(m.group(1)), m.group(2).strip()))
    return out

def collect_tables(text):
    """收集文档里所有含'曲名'的 markdown 表格，返回 [{'headers','rows','sub'}]。
       sub = 表格前最近一个 '### 子标题'（用于把作品归到对应节）。"""
    lines = text.splitlines()
    tables = []
    cur_sub = None
    for i, ln in enumerate(lines):
        if ln.startswith("### "):
            cur_sub = ln[4:].strip()
            continue
        if ln.strip().startswith("|") and "曲名" in ln:
            headers = [c.strip() for c in ln.strip().strip("|").split("|")]
            rows = []
            j = i + 1
            if j < len(lines) and set(lines[j].replace("|", "").strip()) <= set("-: "):
                j += 1
            while j < len(lines) and lines[j].strip().startswith("|"):
                cells = [c.strip() for c in lines[j].strip().strip("|").split("|")]
                if any(cells):
                    rows.append(cells)
                j += 1
            tables.append({"headers": headers, "rows": rows, "sub": cur_sub})
            cur_sub = None
    return tables

def rows_to_works(headers, rows):
    ci = headers.index("曲名") if "曲名" in headers else 0
    auth_cols = [h for h in headers if "作者" in h]
    ai = headers.index(auth_cols[0]) if auth_cols else (headers.index("作者/来源") if "作者/来源" in headers else 1)
    gi = next((k for k in range(len(headers)) if "体裁" in headers[k]), 2)
    si = headers.index("节") if "节" in headers else None
    out = []
    for r in rows:
        name = re.sub(r'[《》]', '', r[ci]).strip()
        author = r[ai].strip() if ai < len(r) else ""
        genre = r[gi].strip() if gi < len(r) else ""
        sec = r[si].strip() if (si is not None and si < len(r)) else None
        out.append({"name": name, "author": author, "genre": genre, "sec": sec})
    return out

def parse_chapter(path):
    text = open(path, encoding="utf-8").read()
    uid = os.path.basename(path)[0:4]  # ch01
    # 单元名
    m = re.search(r'#\s*第[一二三四五六七八九十]+单元\s*([^（(\n]+)', text)
    if not m:
        m = re.search(r'#\s*序篇\s*([^（(\n]+)', text)
    unit_name = m.group(1).strip() if m else uid
    parsed = parse_section_names(text)
    sections = [n for _, n in parsed]
    ordinals = [o for o, _ in parsed]
    if not sections:
        sections = [unit_name]
    # Core Idea
    core = re.search(r'## Core Idea\s*(.*?)(?=\n## )', text, re.S)
    core = core.group(1).strip() if core else ""
    # Frameworks
    frameworks = []
    for blk in re.findall(r'## Frameworks Introduced\s*(.*?)(?=\n## )', text, re.S):
        for item in re.findall(r'-\s*\*\*(.+?)\*\*[：:]\s*(.*?)(?=\n\s*-\s*\*\*|\Z)', blk, re.S):
            term = item[0].strip(); d = re.sub(r'\s+', ' ', item[1].strip())
            frameworks.append({"term": term, "def": d})
    # Key Concepts
    concepts = []
    for blk in re.findall(r'## Key Concepts\s*(.*?)(?=\n## )', text, re.S):
        for item in re.findall(r'-\s*\*\*(.+?)\*\*[：:]\s*(.*?)(?=\n\s*-\s*\*\*|\Z)', blk, re.S):
            concepts.append({"term": item[0].strip(), "def": re.sub(r'\s+', ' ', item[1].strip())})
    # Anti-patterns
    anti = []
    for blk in re.findall(r'## Anti-patterns\s*(.*?)(?=\n## )', text, re.S):
        for item in re.findall(r'-\s*\*\*(.+?)\*\*[：:]\s*(.*?)(?=\n\s*-\s*\*\*|\Z)', blk, re.S):
            anti.append(f"{item[0].strip()}：{re.sub(r'\s+',' ',item[1].strip())}")
    # Works：把作品按「节」/「子标题」/「内容」分到各 section
    sec_works = [[] for _ in sections]
    tables = collect_tables(text)
    global_pool = []
    for t in tables:
        if not t["rows"]:
            continue
        works = rows_to_works(t["headers"], t["rows"])
        has_sec_col = "节" in t["headers"]
        if has_sec_col:  # 有节列（上篇）：按绝对节序号映射
            for w in works:
                msec = re.search(r'第([一二三四五六七八九十]+)节', w["sec"] or "")
                ordv = cn_to_int(msec.group(1)) if msec else None
                idx = ordinals.index(ordv) if (ordv is not None and ordv in ordinals) else None
                if idx is None:
                    idx = 0
                sec_works[idx].append(w)
        elif t["sub"] and any(t["sub"] == s or s in t["sub"] or t["sub"] in s for s in sections):
            # 子标题与节名对应（如 汉族民歌 / 少数民族民歌）
            idx = next((i for i, s in enumerate(sections)
                        if (t["sub"] == s or s in t["sub"] or t["sub"] in s)), 0)
            sec_works[idx].extend(works)
        else:  # 无节列、无对应子标题（下篇常见）：进入全局池，按内容匹配
            global_pool.extend(works)
    # 全局池：按作者/体裁/曲名与节名的关键词重叠度归并
    if global_pool:
        # 少量明确的"人物/板块"归并提示，避免把代表作放错节
        HINTS = {"人民音乐家": ["聂耳", "冼星海"],
                 "学堂乐歌": ["李叔同", "沈心工", "萧友梅", "黄自", "赵元任"]}
        for w in global_pool:
            best, bestscore = 0, -1
            blob = (w["author"] or "") + (w["genre"] or "") + (w["name"] or "")
            for i, sname in enumerate(sections):
                score = 3 if sname in blob else 0
                score += sum(1 for c in sname if c in blob) * 0.1
                if sname in HINTS and any(h in (w["author"] or "") for h in HINTS[sname]):
                    score += 10
                if score > bestscore:
                    bestscore, best = score, i
            sec_works[best].append(w)
    return {
        "uid": uid, "unit_name": unit_name, "group": group_of(uid),
        "sections": sections, "core": core, "frameworks": frameworks,
        "concepts": concepts, "anti": anti, "sec_works": sec_works,
    }

# ---------- 变易教学设计（每单元） ----------
def variation_design(ch):
    concepts = ch["concepts"]; frameworks = ch["frameworks"]
    focal = (concepts[0]["term"] if concepts else (frameworks[0]["term"] if frameworks else "核心概念"))
    wnames = []
    for sw in ch["sec_works"]:
        for w in sw:
            if w["name"] not in wnames: wnames.append(w["name"])
    if not wnames:  # 兜底从单元名
        wnames = [ch["unit_name"]]
    # 审辨焦点
    focus = []
    if ch["core"]:
        focus.append(ch["core"])
    if concepts:
        focus.append(f"辨明「{concepts[0]['term']}」的含义与听辨方法。")
    if len(concepts) >= 2:
        focus.append(f"区分「{concepts[0]['term']}」与「{concepts[1]['term']}」的边界，避免混淆。")
    if not focus:
        focus.append(f"围绕「{focal}」建立清晰辨知。")
    # 四变易图式
    if len(wnames) >= 2:
        contrast = (f"对比《{wnames[0]}》与《{wnames[-1]}》：并置聆听，引导学生聚焦「{focal}」"
                    f"听辨二者处理的不同，体会同一母题如何因要素差异而生出不同音乐性格。")
    else:
        contrast = f"对比本单元内最相近的两首作品，聚焦「{focal}」听辨其处理差异。"
    generalize = (f"类合：将《{'》《'.join(wnames[:3])}》并听，归纳它们在「{focal}」上的共同表达，"
                 f"形成对这类音乐的概括性认识。")
    separate = (f"区分：暂时悬置其他要素，单点聚焦「{focal}」，让学生只听这一层，"
                f"建立清晰辨知后再放回整体。")
    fuse = (f"融合：设计聆听任务，让学生同时关照「{focal}」与情感内涵"
            f"（如边听边标情绪 / 写联想），在综合中深化理解。")
    return {"focus": focus, "moves": [("对比", contrast), ("类合", generalize),
                                      ("区分", separate), ("融合", fuse)]}

# ---------- PPT 构建 ----------
def add_textbox(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    return tb

def style_run(r, size, color=INK, bold=False, italic=False):
    r.font.size = Pt(size); r.font.color.rgb = color
    r.font.bold = bold; r.font.italic = italic
    r.font.name = "Microsoft YaHei"

def slide_blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def header(slide, title, sub=None):
    tb = add_textbox(slide, 0.6, 0.35, 12.1, 0.9)
    p = tb.text_frame.paragraphs[0]
    style_run(p.add_run(), 26, BRAND, bold=True)
    p.runs[0].text = title
    if sub:
        p2 = tb.text_frame.add_paragraph()
        style_run(p2.add_run(), 14, MUTED)
        p2.runs[0].text = sub

def bullets(slide, items, l=0.7, t=1.4, w=12.0, h=5.6, size=16, gap=6):
    tb = add_textbox(slide, l, t, w, h)
    tf = tb.text_frame
    for i, it in enumerate(items):
        if isinstance(it, tuple):
            txt, lvl, bold = it
        else:
            txt, lvl, bold = it, 0, False
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        p.space_after = Pt(gap)
        style_run(p.add_run(), size if lvl == 0 else size - 2, INK if lvl == 0 else MUTED, bold)
        p.runs[0].text = ("• " if lvl == 0 else "– ") + txt
    return tb

def build_ppt(ch, sec_idx, quiz):
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    sec_name = ch["sections"][sec_idx]
    sec_works = ch["sec_works"][sec_idx]
    var = variation_design(ch)

    # 1. 封面
    s = slide_blank(prs)
    tb = add_textbox(s, 0.7, 1.6, 11.5, 1.3)
    p = tb.text_frame.paragraphs[0]
    style_run(p.add_run(), 30, BRAND2, bold=True)
    if ch["group"] == "序篇":
        p.runs[0].text = ch["unit_name"]
    else:
        p.runs[0].text = f"第{ch['uid'][2:]}单元　{ch['unit_name']}"
    tb2 = add_textbox(s, 0.7, 2.9, 11.5, 0.9)
    p2 = tb2.text_frame.paragraphs[0]
    style_run(p2.add_run(), 22, INK)
    p2.runs[0].text = f"第{sec_idx+1}节　{sec_name}"
    tb3 = add_textbox(s, 0.7, 4.0, 11.5, 0.6)
    p3 = tb3.text_frame.paragraphs[0]
    style_run(p3.add_run(), 15, MUTED)
    p3.runs[0].text = f"河北峰峰第一中学 · 音乐鉴赏（人音版必修 2019）· {ch['group']}"
    if os.path.exists(LOGO):
        s.shapes.add_picture(LOGO, Inches(11.5), Inches(0.45), width=Inches(1.15))
    if os.path.exists(NAMEIMG):
        s.shapes.add_picture(NAMEIMG, Inches(0.7), Inches(4.7), width=Inches(3.0))

    # 2. 学习目标 / Core Idea
    s = slide_blank(prs); header(s, "学习目标", "本课要带学生走到哪里")
    bullets(s, [ch["core"]] if ch["core"] else ["（见章节正文）"], size=18)

    # 3. 核心框架 / 概念
    s = slide_blank(prs); header(s, "核心框架 / 概念", "鉴赏的工具")
    items = []
    for f in ch["frameworks"][:6]:
        items.append((f"{f['term']}：{f['def'][:60]}{'…' if len(f['def'])>60 else ''}", 0, False))
    for c in ch["concepts"][:6]:
        items.append((f"{c['term']}：{c['def'][:60]}{'…' if len(c['def'])>60 else ''}", 1, False))
    if not items: items = [("（本单元以聆听体验为主，详见教材）", 0, False)]
    bullets(s, items, size=15)

    # 4. 易错提醒
    if ch["anti"]:
        s = slide_blank(prs); header(s, "易错提醒", "学生最容易踩的坑")
        bullets(s, ch["anti"][:4], size=16)

    # 5. 作品鉴赏
    s = slide_blank(prs); header(s, "作品鉴赏", "本课曲目与聆听要点")
    if sec_works:
        rows = len(sec_works) + 1
        tbl = s.shapes.add_table(rows, 3, Inches(0.7), Inches(1.5), Inches(11.9), Inches(0.5*rows)).table
        tbl.columns[0].width = Inches(3.4); tbl.columns[1].width = Inches(5.0); tbl.columns[2].width = Inches(3.5)
        for j, htxt in enumerate(["曲名", "作者 / 来源", "体裁"]):
            cell = tbl.cell(0, j); cell.text = htxt
            style_run(cell.text_frame.paragraphs[0].runs[0] if cell.text_frame.paragraphs[0].runs else cell.text_frame.paragraphs[0].add_run(), 15, BRAND2, bold=True)
            cell.text_frame.paragraphs[0].runs[0].text = htxt
            cell.fill.solid(); cell.fill.fore_color.rgb = PANEL
        for i, w in enumerate(sec_works, 1):
            for j, val in enumerate([w["name"], w["author"], w["genre"]]):
                cell = tbl.cell(i, j); cell.text = val
                style_run(cell.text_frame.paragraphs[0].runs[0] if cell.text_frame.paragraphs[0].runs else cell.text_frame.paragraphs[0].add_run(), 13, INK)
                cell.text_frame.paragraphs[0].runs[0].text = val
    else:
        bullets(s, ["（本课本节无单独曲目，参见单元其他节 / 教材聆听示例）"], size=16)

    # 6. 变易教学设计 - 审辨焦点
    s = slide_blank(prs); header(s, "变易教学设计 · 审辨焦点", "先让学生‘看出来’什么")
    bullets(s, var["focus"], size=17)

    # 7. 变易教学设计 - 四图式
    s = slide_blank(prs); header(s, "变易教学设计 · 四变易图式", "对比 / 类合 / 区分 / 融合")
    items = [(f"【{name}】{desc}", 0, False) for name, desc in var["moves"]]
    bullets(s, items, size=15)

    # 8. 随堂练习
    if quiz:
        for q in quiz:
            s = slide_blank(prs)
            header(s, "随堂练习", f"{q['type'] if 'type' in q else '选择题'}")
            tb = add_textbox(s, 0.7, 1.5, 12.0, 1.2)
            p = tb.text_frame.paragraphs[0]; style_run(p.add_run(), 18, INK, bold=True)
            p.runs[0].text = q["q"]
            opts = q.get("options", [])
            ans = q.get("answer", 0)
            letters = "ABCD"
            oitems = []
            for k, o in enumerate(opts):
                mark = " ✅" if k == ans else ""
                oitems.append((f"{letters[k]}. {o}{mark}", 0, (k == ans)))
            bullets(s, oitems, t=2.7, size=16, gap=8)
            if q.get("explain"):
                tb2 = add_textbox(s, 0.7, 6.2, 12.0, 0.8)
                p2 = tb2.text_frame.paragraphs[0]; style_run(p2.add_run(), 12, MUTED, italic=True)
                p2.runs[0].text = "解析：" + q["explain"][:120]
    else:
        s = slide_blank(prs); header(s, "随堂练习", "暂无自动题目")
        bullets(s, ["（可在网页 index.html 对应单元随堂练习中选用）"], size=16)

    # 9. 配套资源
    s = slide_blank(prs); header(s, "配套资源", "与网页电子课本联动")
    bullets(s, [
        "打开 index.html：本单元重点知识讲解 + 即时判分练习 + 术语速查。",
        "右下角 📝 便签：上课随手记校本补充，仅存本机浏览器。",
        "校徽与校名已置入页面，投影即见学校标识。",
    ], size=16)

    fname = f"{ch['uid']}-{sec_idx+1} {ch['unit_name']}-{sec_name}.pptx"
    fname = fname.replace("/", "·").replace("\\", "·")
    path = os.path.join(OUT, fname)
    prs.save(path)
    return path

# ---------- 主流程 ----------
def load_quiz():
    txt = open(os.path.join(ROOT, "data.js"), encoding="utf-8").read()
    js = txt.split("window.COURSE_DATA = ", 1)[1]
    js = js.rstrip().rstrip(";")
    data = json.loads(js)
    return {u["id"]: u.get("quiz", []) for u in data["units"]}

def main():
    quiz_map = load_quiz()
    files = sorted(glob.glob(os.path.join(SKILL, "chapters", "ch*.md")))
    total = 0
    for path in files:
        ch = parse_chapter(path)
        qall = quiz_map.get(ch["uid"], [])
        nsec = len(ch["sections"])
        # 随堂练习按节轮询分配
        for si in range(nsec):
            qsec = [qall[k] for k in range(len(qall)) if k % nsec == si]
            p = build_ppt(ch, si, qsec)
            total += 1
            print(f"  {os.path.basename(p)}  ({len(qsec)}题)")
    print(f"共生成 {total} 个 PPT -> {OUT}")

if __name__ == "__main__":
    main()
